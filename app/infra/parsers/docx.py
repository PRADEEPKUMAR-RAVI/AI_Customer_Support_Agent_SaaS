"""DOCX parser — extracts paragraph + table text as a single segment (docx has no pages)."""

from __future__ import annotations

import io

from app.infra.parsers.base import ParseResult, UnsupportedFormat

DOCX_EXTENSIONS = {".docx"}
DOCX_CONTENT_TYPES = {
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
}


class DocxParser:
    def parse(self, data: bytes) -> list[ParseResult]:
        from docx import Document  # lazy — heavy import

        try:
            document = Document(io.BytesIO(data))
        except Exception as exc:  # noqa: BLE001 — corrupt/unsupported -> terminal failure
            raise UnsupportedFormat(f"could not read DOCX: {exc}") from exc

        parts = [p.text for p in document.paragraphs if p.text.strip()]
        for table in document.tables:
            for row in table.rows:
                cells = [c.text.strip() for c in row.cells]
                if any(cells):
                    parts.append(" | ".join(cells))  # keep table rows readable + retrievable
        text = "\n\n".join(parts).strip()
        if not text:
            raise UnsupportedFormat("DOCX has no extractable text")
        return [ParseResult(text=text)]


def can_parse_docx(filename: str | None, content_type: str | None) -> bool:
    ext = ("." + filename.rsplit(".", 1)[-1].lower()) if filename and "." in filename else ""
    ct = (content_type or "").split(";", 1)[0].strip().lower()
    return ext in DOCX_EXTENSIONS or ct in DOCX_CONTENT_TYPES
